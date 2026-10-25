# -*- coding: utf-8 -*-
"""Собирает content/<курс>/content.xlsx из авторского источника (content_src.py)
и существующего facts.json, затем вызывает xlsx2json.py.
Запуск из корня репозитория: python3 tools/build_content.py"""
import json, re, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, os.path.dirname(__file__))
import content_src as S

COURSE = "srednie-veka-6"
ROOT = os.path.join(os.path.dirname(__file__), "..", "content", COURSE)

ROMAN = {"XV в.": 15, "XIV в.": 14, "XIII в.": 13}

def year_of(d):
    if d in ROMAN: return (ROMAN[d] - 1) * 100 + 50
    m = re.search(r"\d{3,4}", d)
    return int(m.group()) if m else None

def century(y): return (y - 1) // 100 + 1

facts = json.load(open(os.path.join(ROOT, "facts.json"), encoding="utf-8"))
dates = [f for f in facts if f["type"] == "дата"]

wb = Workbook()
H = Font(name="Arial", bold=True, color="FFFFFF", size=10)
F = PatternFill("solid", fgColor="1F3FBF")
B = Font(name="Arial", size=10)

def sheet(name, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in rows: ws.append(list(r))
    for c in ws[1]: c.font, c.fill = H, F
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = B; c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate(widths, 1): ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    return ws

wb.remove(wb.active)
sheet("Легенда", ["Лист", "Что внутри"], [
    ("Даты", "Одна строка — одна дата. assoc — ассоциация для запоминания (показывается при ошибке), more — мини-факт за значком «хочу все знать», hook — шутка при верном ответе"),
    ("Термины", "Термин и определение. Дистракторы берутся из определений других терминов"),
    ("Персоналии", "Имя и кто это"),
    ("Причины", "Причина → верное следствие и два неверных"),
    ("Тексты", "Текст с пропусками {0} {1} {2}; ответы и лишние варианты через |"),
    ("Правда-ложь", "Утверждение, верно ли, пояснение для ложных"),
    ("Уроки", "Разбивка дат по урокам (id через запятую). Урок = 8 заданий"),
    ("Главы", "Какие параграфы в какой главе — для режима «Контрольная»"),
    ("Правило", "Правь любой лист, потом python3 tools/xlsx2json.py — и JSON обновится. Столбец id не менять"),
], [16, 120])

sheet("Даты", ["id", "topic", "date", "year", "century", "event", "person", "term", "hook", "assoc_type", "assoc", "more", "check"],
      [(f["id"], f["topic"], f["date"], year_of(f["date"]), century(year_of(f["date"])), f["event"], f["person"], f["term"], f["hook"],
        S.ASSOC.get(f["id"], ("", ""))[0], S.ASSOC.get(f["id"], ("", ""))[1], S.MORE.get(f["id"], ""), f["check"]) for f in dates],
      [5, 8, 13, 7, 8, 50, 22, 18, 50, 10, 60, 50, 24])
sheet("Термины", ["id", "topic", "term", "definition"], [(i + 1, *t) for i, t in enumerate(S.TERMS)], [5, 8, 26, 80])
sheet("Персоналии", ["id", "topic", "name", "who"], [(i + 1, *p) for i, p in enumerate(S.PERSONS)], [5, 8, 28, 80])
sheet("Причины", ["id", "topic", "cause", "effect", "wrong1", "wrong2"], [(i + 1, t, c, e, w[0], w[1]) for i, (t, c, e, w) in enumerate(S.CAUSES)], [5, 8, 50, 50, 36, 36])
sheet("Тексты", ["id", "topic", "text", "answers", "extra"], [(i + 1, t, x, "|".join(a), "|".join(e)) for i, (t, x, a, e) in enumerate(S.CLOZE)], [5, 8, 90, 30, 30])
sheet("Правда-ложь", ["id", "topic", "statement", "is_true", "note"], [(i + 1, t, s, "да" if v else "нет", n) for i, (t, s, v, n) in enumerate(S.TRUEFALSE)], [5, 8, 60, 8, 60])
sheet("Уроки", ["id", "topic", "title", "fact_ids"], [(i + 1, t, n, ",".join(map(str, ids))) for i, (t, n, ids) in enumerate(S.LESSONS)], [5, 8, 44, 30])
sheet("Главы", ["chapter", "title", "topics"], [(k, v["title"], ",".join(v["topics"])) for k, v in S.CHAPTERS.items()], [8, 50, 30])

out = os.path.join(ROOT, "content.xlsx")
wb.save(out)
print("xlsx:", out)
os.system(f"{sys.executable} {os.path.join(os.path.dirname(__file__), 'xlsx2json.py')}")
