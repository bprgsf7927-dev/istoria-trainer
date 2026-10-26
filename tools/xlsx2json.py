# -*- coding: utf-8 -*-
"""content/<курс>/content.xlsx → JSON-файлы рядом. Источник правды — xlsx.
Запуск из корня репозитория: python3 tools/xlsx2json.py"""
import json, os
from openpyxl import load_workbook

COURSE = "srednie-veka-6"
ROOT = os.path.join(os.path.dirname(__file__), "..", "content", COURSE)
wb = load_workbook(os.path.join(ROOT, "content.xlsx"), data_only=True)

def rows(name):
    ws = wb[name]
    hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None: continue
        yield {h: ("" if v is None else v) for h, v in zip(hdr, r)}

def dump(name, data):
    with open(os.path.join(ROOT, name), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(name, len(data))

facts = []
for r in rows("Даты"):
    r["id"] = int(r["id"]); r["year"] = int(r["year"]); r["century"] = int(r["century"])
    r["topic"] = str(r["topic"]); r["date"] = str(r["date"])
    facts.append(r)
dump("facts.json", facts)

dump("terms.json", [{"id": int(r["id"]), "topic": str(r["topic"]), "term": r["term"], "definition": r["definition"]} for r in rows("Термины")])
dump("persons.json", [{"id": int(r["id"]), "topic": str(r["topic"]), "name": r["name"], "who": r["who"]} for r in rows("Персоналии")])
dump("causes.json", [{"id": int(r["id"]), "topic": str(r["topic"]), "cause": r["cause"], "effect": r["effect"], "wrong": [r["wrong1"], r["wrong2"]]} for r in rows("Причины")])
dump("cloze.json", [{"id": int(r["id"]), "topic": str(r["topic"]), "text": r["text"], "answers": r["answers"].split("|"), "extra": r["extra"].split("|")} for r in rows("Тексты")])
dump("truefalse.json", [{"id": int(r["id"]), "topic": str(r["topic"]), "statement": r["statement"], "isTrue": r["is_true"] == "да", "note": r["note"]} for r in rows("Правда-ложь")])
dump("lessons.json", [{"id": int(r["id"]), "topic": str(r["topic"]), "title": r["title"], "factIds": [int(x) for x in str(r["fact_ids"]).split(",")]} for r in rows("Уроки")])

topics = {"course": "История Средних веков, 6 класс (Мединский, Чубарьян, 2025)",
          "topics": [{"id": str(r["id"]), "paragraph": r["paragraph"], "chapter": r["chapter"], "title": r["title"], "block": r["block"]} for r in rows("Параграфы")]}
topics["chapters"] = [{"id": r["chapter"], "title": r["title"], "topics": r["topics"].split(",")} for r in rows("Главы")]
with open(os.path.join(ROOT, "topics.json"), "w", encoding="utf-8") as f:
    json.dump(topics, f, ensure_ascii=False, indent=1)
print("topics.json + chapters")
